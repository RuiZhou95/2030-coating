#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Validate the standalone coating Supporting Information artifact."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile

import pandas as pd
from docx import Document

from project_paths import PROJECT_ROOT as PROJECT

DOCX = PROJECT / "supplementary/Supplementary_Information_Coating_CN.docx"
OUT = PROJECT / "supplementary/supplementary_validation.json"


def sha256(path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main():
    checks = []

    def check(name, passed, detail):
        checks.append({"name": name, "passed": bool(passed), "detail": str(detail)})

    check("docx_exists", DOCX.exists(), DOCX)
    if not DOCX.exists():
        OUT.write_text(json.dumps({"status": "FAIL", "checks": checks}, ensure_ascii=False, indent=2), encoding="utf-8")
        raise SystemExit(1)

    doc = Document(DOCX)
    text = "\n".join(p.text for p in doc.paragraphs)
    table_text = "\n".join(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
    all_text = text + "\n" + table_text

    check("docx_size", DOCX.stat().st_size > 1_500_000, DOCX.stat().st_size)
    check("table_count", len(doc.tables) == 7, len(doc.tables))
    check("image_count", len(doc.inline_shapes) == 4, len(doc.inline_shapes))
    expected_rows = [6, 6, 15, 11, 11, 6, 7]
    actual_rows = [len(t.rows) for t in doc.tables]
    check("table_row_counts", actual_rows == expected_rows, actual_rows)

    required = [
        "S1 两类功能涂层", "S6 负对照", "图S1", "图S4", "表S1", "表S7",
        "salt_spray_pass_h", "salt_spray_time_h", "72", "120", "71", "6×5×4",
        "φF", "ψ", "T116", "机器学习", "Rui Zhou", "Luyao Bao",
        "Meirong Cai", "baoluyao@licp.cas.cn", "caimr@licp.cas.cn",
        "GB/T 5210", "GB/T 1771", "500 g载荷和1000 r", "1000 g载荷和1000 r",
    ]
    missing = [term for term in required if term not in all_text]
    check("required_terms", not missing, missing or "all present")
    forbidden = [
        "[[FIG", "[[TABLE", "[[PAGEBREAK]]", "TODO", "TBD", "物理坐标", "物理设计坐标",
        "物理化学假设", "材料假设", "假设", "不构成", "仍需", "不宜", "仅", "25/4", "35/4", "15/4",
        "99.99-100.02", "材料平台", "功能涂层平台", "窗口",
    ]
    found_forbidden = [term for term in forbidden if term in all_text]
    check("no_internal_markers", not found_forbidden, found_forbidden or "none")
    check("materials_first_language", "材料图谱" not in all_text and "性能图谱" not in all_text, "no atlas wording")
    check("no_author_placeholder", "同主文" not in all_text and "***" not in all_text, "actual English author block")
    author_paragraph = next((p for p in doc.paragraphs if p.text.startswith("Rui Zhoua, Luyao Baoa*, Meirong Caia*")), None)
    author_markers = [run for run in author_paragraph.runs if run.text in {"a", "a*"}] if author_paragraph is not None else []
    author_superscript_ok = len(author_markers) == 3 and all(run.font.superscript is True for run in author_markers)
    check("author_affiliations_superscript", author_superscript_ok, [run.text for run in author_markers])
    check("salt_time_explanatory_only", all_text.count("salt_spray_time_h") == 1, all_text.count("salt_spray_time_h"))
    salt_definition_ok = all(term in all_text for term in [
        "耐腐蚀性能采用中性盐雾试验评价",
        "所得累计时间定义为盐雾寿命",
        "原始工作簿中记录为salt_spray_pass_h",
    ])
    check("salt_spray_definition_self_contained", salt_definition_ok, "method, endpoint, lifetime definition and source-field mapping")
    reproduction_phrases = ["数据处理与复现信息", "原始工作簿保持只读", "固定脚本生成", "图表均可追溯"]
    found_reproduction_phrases = [term for term in reproduction_phrases if term in all_text]
    check("reproduction_note_removed", not found_reproduction_phrases, found_reproduction_phrases or "removed")
    priority_explained = all(term in all_text for term in [
        "偏向不利方向的保守估计",
        "实验复核优先级综合实测性能",
        "排位越靠前，越适合优先复核",
    ]) and "保守盐雾" not in all_text and "保守复核顺序" not in all_text
    check("review_priority_and_conservative_values_explained", priority_explained, "adverse-direction estimate and experimental review priority defined")
    formulation_family_explained = all(term in all_text for term in [
        "具有相同主要组分水平与工艺组合的样本归为同一配方类别",
        "formulation_family作为类别标签",
        "同类配方不会同时分布在训练集和测试集中",
    ])
    check("formulation_family_explained", formulation_family_explained, "grouping label and holdout logic defined")
    ilr_explained = all(term in all_text for term in [
        "等距对数比（isometric log-ratio, ILR）",
        "转换为相互独立的对数比坐标",
        "检验模型结论是否依赖直接质量分数表示",
    ])
    check("ilr_explained", ilr_explained, "full name, coordinate meaning and purpose defined")
    check("random_seed_sentence_removed", "随机种子固定为20260823" not in all_text, "removed")
    crossrefs = {label: all_text.count(label) for label in ["图S1", "图S2", "图S3", "图S4", "表S1", "表S2", "表S3", "表S4", "表S5", "表S6", "表S7"]}
    check("all_displays_cited", all(count >= 2 for count in crossrefs.values()), crossrefs)
    excessive_decimals = sorted(set(re.findall(r"(?<![A-Za-z_])[-+]?\d+\.\d{3,}", all_text)))
    check("all_numeric_decimals_le_two", not excessive_decimals, excessive_decimals or "all numeric decimals <= 2")
    variance_table = next((table for table in doc.tables if table.cell(0, 0).text == "响应" and table.cell(0, 2).text == "观测平方和"), None)
    variance_values = [row.cells[2].text for row in variance_table.rows[1:]] if variance_table is not None else []
    variance_plain = bool(variance_values) and variance_values[0] == "35770000" and all(not re.search(r"[eE][+-]?\d+", value) for value in variance_values)
    check("variance_sums_plain_decimal", variance_plain, variance_values)

    aliases = pd.read_csv(PROJECT / "data/removed_target_aliases.csv", encoding="utf-8-sig")
    alias_ok = (
        len(aliases) == 2
        and set(aliases["column"]) == {"salt_spray_time_h"}
        and set(aliases["duplicate_of"]) == {"salt_spray_pass_h"}
        and aliases["action"].str.contains("dropped", case=False).all()
    )
    check("target_alias_removed", alias_ok, aliases.to_dict(orient="records"))

    split = json.loads((PROJECT / "data/split_manifest.json").read_text(encoding="utf-8"))
    split_ok = (
        split["frozen_before_modeling"] is True
        and split["primer"]["holdout_n"] == 8
        and split["primer"]["training_n"] == 64
        and split["topcoat"]["holdout_n"] == 8
        and split["topcoat"]["training_n"] == 112
    )
    check("frozen_split", split_ok, {"primer": split["primer"], "topcoat": split["topcoat"]})

    variance = pd.read_csv(PROJECT / "results/topcoat_variance_decomposition.csv", encoding="utf-8-sig")
    sums = variance.groupby("target")["variance_pct"].sum().round(8).to_dict()
    check("variance_sums_100", all(abs(v - 100.0) < 1e-6 for v in sums.values()), sums)

    uncertainty = pd.read_csv(PROJECT / "results/uncertainty_summary.csv", encoding="utf-8-sig")
    frozen = uncertainty[uncertainty["set"] == "frozen_holdout"]
    unc_ok = (
        len(frozen) == 5
        and abs(frozen["bootstrap_90_coverage"].min() - 0.125) < 1e-9
        and abs(frozen["bootstrap_90_coverage"].max() - 0.5) < 1e-9
        and abs(frozen["conformal_90_coverage"].min() - 0.75) < 1e-9
        and abs(frozen["conformal_90_coverage"].max() - 1.0) < 1e-9
    )
    check("uncertainty_bounds", unc_ok, frozen[["layer", "target", "bootstrap_90_coverage", "conformal_90_coverage"]].to_dict(orient="records"))

    primer = pd.read_csv(PROJECT / "results/robust_candidates_primer.csv", encoding="utf-8-sig").head(10)
    topcoat = pd.read_csv(PROJECT / "results/robust_candidates_topcoat.csv", encoding="utf-8-sig").head(10)
    candidate_ok = (
        str(primer.iloc[0]["sample_id"]).endswith("08")
        and str(primer.iloc[1]["sample_id"]).endswith("42")
        and str(topcoat.iloc[0]["sample_id"]).endswith("112")
        and "P08" in table_text and "P42" in table_text and "T112" in table_text
    )
    check("candidate_tables_traceable", candidate_ok, {"primer_first": primer.iloc[0]["sample_id"], "topcoat_first": topcoat.iloc[0]["sample_id"]})

    with zipfile.ZipFile(DOCX) as zf:
        names = set(zf.namelist())
        document_xml = zf.read("word/document.xml").decode("utf-8", errors="replace")
    check("no_tracked_changes", "<w:ins" not in document_xml and "<w:del" not in document_xml, "document.xml")
    check("no_comments_part", "word/comments.xml" not in names, "word/comments.xml" in names)
    nonblack = []
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.color.type is not None and run.font.color.rgb is not None and str(run.font.color.rgb).upper() != "000000":
                nonblack.append(str(run.font.color.rgb))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.color.type is not None and run.font.color.rgb is not None and str(run.font.color.rgb).upper() != "000000":
                            nonblack.append(str(run.font.color.rgb))
    check("all_text_black", not nonblack, sorted(set(nonblack)) or "all black")

    failures = [c for c in checks if not c["passed"]]
    result = {
        "status": "PASS" if not failures else "FAIL",
        "failed_count": len(failures),
        "docx_sha256": sha256(DOCX),
        "docx_bytes": DOCX.stat().st_size,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "images": len(doc.inline_shapes),
        "checks": checks,
    }
    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if not failures else 1)


if __name__ == "__main__":
    main()

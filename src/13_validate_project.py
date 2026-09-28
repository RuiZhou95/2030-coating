#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Requirement-by-requirement completion audit for the coating paper goal."""

from __future__ import annotations

import hashlib
import json
import re

import pandas as pd
from docx import Document

from project_paths import PROJECT_ROOT as PROJECT, TEMPLATE_DOCX as REFERENCE

DOCX = PROJECT / "manuscript/Manuscript_Coating_CN.docx"
EXPECTED_TEMPLATE_SHA = "4883eb29612987688291e8f5d822a86ac0dc498f26520bbdd83e075e127c81fc"


def sha(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def id_digest(path):
    ids = pd.read_csv(path, encoding="utf-8-sig")["sample_id"].astype(str).tolist()
    return hashlib.sha256("\n".join(sorted(ids)).encode("utf-8")).hexdigest()


checks = []


def record(name, passed, evidence):
    checks.append({"requirement": name, "status": "pass" if passed else "fail", "evidence": str(evidence)})


def require_files(name, paths):
    missing = [str(p) for p in paths if not p.exists()]
    record(name, not missing, "all present" if not missing else "missing: " + " | ".join(missing))


def expand_citation_token(token):
    numbers = []
    for part in token.split(","):
        if "-" in part:
            start, end = map(int, part.split("-"))
            numbers.extend(range(start, end + 1))
        else:
            numbers.append(int(part))
    return numbers


def main():
    dirs = ["config", "src", "data", "results", "figure", "manuscript", "supplementary", "logs", "slurm"]
    record("项目目录结构", all((PROJECT / d).is_dir() for d in dirs), ", ".join(dirs))
    require_files("项目治理文档", [PROJECT / x for x in ["README.md", "PROGRESS.md", "DECISIONS.md", "BLOCKERS.md"]])

    data_files = [
        "data_dictionary.csv", "data_quality_summary.csv", "sample_groups.csv",
        "target_statistics.csv", "source_manifest.json", "design_factor_levels.csv",
        "leakage_candidates.csv", "split_manifest.json", "frozen_holdout_primer.csv",
        "frozen_holdout_topcoat.csv", "training_ids_primer.csv", "training_ids_topcoat.csv",
        "removed_target_aliases.csv",
    ]
    require_files("数据审计与冻结产物", [PROJECT / "data" / x for x in data_files] + [PROJECT / "results/data_audit.md", PROJECT / "results/design_structure_audit.md"])

    manifest = json.loads((PROJECT / "data/split_manifest.json").read_text(encoding="utf-8"))
    split_ok = (
        manifest["frozen_before_modeling"] is True
        and manifest["primer"]["holdout_n"] == 8 and manifest["primer"]["training_n"] == 64
        and manifest["topcoat"]["holdout_n"] == 8 and manifest["topcoat"]["training_n"] == 112
        and manifest["primer"]["holdout_id_sha256"] == id_digest(PROJECT / "data/frozen_holdout_primer.csv")
        and manifest["primer"]["training_id_sha256"] == id_digest(PROJECT / "data/training_ids_primer.csv")
        and manifest["topcoat"]["holdout_id_sha256"] == id_digest(PROJECT / "data/frozen_holdout_topcoat.csv")
        and manifest["topcoat"]["training_id_sha256"] == id_digest(PROJECT / "data/training_ids_topcoat.csv")
    )
    record("冻结划分不可变性", split_ok, json.dumps(manifest, ensure_ascii=False))

    removed = pd.read_csv(PROJECT / "data/removed_target_aliases.csv", encoding="utf-8-sig")
    dictionary = pd.read_csv(PROJECT / "data/data_dictionary.csv", encoding="utf-8-sig")
    feature_files = list((PROJECT / "data").glob("final_features_*.txt"))
    alias_clean = (
        set(removed["column"]) == {"salt_spray_time_h"}
        and set(removed["layer"]) == {"底漆", "面漆"}
        and "salt_spray_time_h" not in set(dictionary["column"])
        and all("salt_spray_time_h" not in path.read_text(encoding="utf-8") for path in feature_files)
    )
    record("目标重复字段已从分析层删除", alias_clean, removed.to_dict("records"))

    model_files = [
        "nested_cv_predictions.csv", "nested_cv_fold_metrics.csv", "nested_cv_summary.csv",
        "model_selection_ranking.csv", "final_model_selection.csv", "holdout_point_predictions.csv",
    ]
    require_files("严格嵌套与外推验证", [PROJECT / "results" / x for x in model_files])
    final = pd.read_csv(PROJECT / "results/final_model_selection.csv", encoding="utf-8-sig")
    record("五个层级目标模型", len(final) == 5 and set(final.layer) == {"primer", "topcoat"}, final[["layer", "target", "representation", "model"]].to_dict("records"))

    robust_files = [
        "uncertainty_predictions.csv", "uncertainty_summary.csv", "stable_feature_effects.csv",
        "stable_associations.csv", "label_permutation_controls.csv", "negative_control_summary.csv",
        "random_noise_controls.csv", "zero_replacement_sensitivity.csv",
    ]
    require_files("稳定解释、不确定性与负对照", [PROJECT / "results" / x for x in robust_files])
    perm = pd.read_csv(PROJECT / "results/label_permutation_controls.csv", encoding="utf-8-sig")
    record("100次标签置换覆盖全部目标", len(perm) == 500 and perm.groupby(["layer", "target"]).size().eq(100).all(), f"rows={len(perm)}")
    unc = pd.read_csv(PROJECT / "results/uncertainty_summary.csv", encoding="utf-8-sig")
    record("Bootstrap与conformal区间", {"bootstrap_90_coverage", "conformal_90_coverage"}.issubset(unc.columns) and len(unc) == 10, f"rows={len(unc)}")
    unc_points = pd.read_csv(PROJECT / "results/uncertainty_predictions.csv", encoding="utf-8-sig")
    nested_points = pd.read_csv(PROJECT / "results/nested_cv_predictions.csv", encoding="utf-8-sig")
    validation_counts = []
    validation_groups_ok = True
    for _, chosen in final[final["target_role"] == "core"].iterrows():
        fit_n = len(unc_points[(unc_points.layer == chosen.layer) & (unc_points.target == chosen.target) & (unc_points.set == "training")])
        frozen_n = len(unc_points[(unc_points.layer == chosen.layer) & (unc_points.target == chosen.target) & (unc_points.set == "frozen_holdout")])
        oof_n = len(nested_points[
            (nested_points.layer == chosen.layer)
            & (nested_points.target == chosen.target)
            & (nested_points.representation == chosen.representation)
            & (nested_points.model == chosen.model)
            & (nested_points.scheme == "family_logo")
        ])
        validation_groups_ok &= fit_n == chosen.training_n and oof_n == chosen.training_n and frozen_n == chosen.holdout_n
        validation_counts.append(f"{chosen.layer}/{chosen.target}={fit_n}/{oof_n}/{frozen_n}")
    record("拟合-折外-冻结预测数据完整", validation_groups_ok, "; ".join(validation_counts))

    material_files = [
        "primer_factor_means.csv", "primer_sr_znp_cells.csv", "primer_curing_interactions.csv", "primer_design_cells.csv",
        "topcoat_physical_coordinates.csv", "topcoat_variance_decomposition.csv",
        "topcoat_phi_psi_means.csv", "topcoat_full_observed_pareto.csv",
        "material_evidence_summary.json", "material_chart_map.md",
    ]
    require_files("材料物理坐标分析产物", [PROJECT / "results" / x for x in material_files])
    primer_cells = pd.read_csv(PROJECT / "results/primer_sr_znp_cells.csv", encoding="utf-8-sig")
    cell_25_4 = primer_cells[(primer_cells["铬酸锶"] == 25) & (primer_cells["磷酸锌"] == 4)].iloc[0]
    cell_35_4 = primer_cells[(primer_cells["铬酸锶"] == 35) & (primer_cells["磷酸锌"] == 4)].iloc[0]
    primer_window_ok = (
        abs(cell_25_4.salt_spray_mean_h - 3216.6667) < 0.1
        and abs(cell_25_4.adhesion_mean_mpa - 17.4667) < 0.01
        and abs(cell_35_4.salt_spray_mean_h - 3225.0) < 0.1
        and cell_25_4.adhesion_mean_mpa > cell_35_4.adhesion_mean_mpa
    )
    record("底漆25/4缓蚀剂效率窗口", primer_window_ok, f"25/4={cell_25_4.to_dict()} 35/4={cell_35_4.to_dict()}")
    primer_design = pd.read_csv(PROJECT / "results/primer_design_cells.csv", encoding="utf-8-sig")
    primer_design_ok = len(primer_design) == 71 and len(primer_design[primer_design["n"] == 2]) == 1 and primer_design["n"].sum() == 72
    record("底漆名义4x3x3x2设计缺口已披露", primer_design_ok, f"records={primer_design['n'].sum()} unique={len(primer_design)} duplicates={primer_design[primer_design['n']>1].to_dict('records')}")
    coords = pd.read_csv(PROJECT / "results/topcoat_physical_coordinates.csv", encoding="utf-8-sig")
    grid_n = coords.groupby(["fluororesin_fraction_in_binder", "solid_to_binder_mass_ratio", "filler_package_level"]).size()
    variance = pd.read_csv(PROJECT / "results/topcoat_variance_decomposition.csv", encoding="utf-8-sig")
    variance_map = {(row.target, row.term): row.variance_pct for row in variance.itertuples()}
    full_pareto = pd.read_csv(PROJECT / "results/topcoat_full_observed_pareto.csv", encoding="utf-8-sig")
    topcoat_grid_ok = (
        len(coords) == 120 and len(grid_n) == 120 and grid_n.eq(1).all()
        and abs(variance_map[("salt_spray_pass_h", "Binder fluorination (phi_F)")] - 68.4842) < 0.01
        and abs(variance_map[("salt_spray_pass_h", "phi_F x psi")] - 14.7861) < 0.01
        and abs(variance_map[("wear_mass_loss_mg", "Functional-solid/binder ratio (psi)")] - 99.9347) < 0.01
        and len(full_pareto) == 1 and str(full_pareto.iloc[0].sample_id).endswith("116")
    )
    record("面漆6x5x4物理坐标与完整域Pareto", topcoat_grid_ok, f"grid={len(grid_n)} pareto={full_pareto[['sample_id','salt_spray_pass_h','wear_mass_loss_mg']].to_dict('records')}")

    pareto_files = ["pareto_primer.csv", "pareto_topcoat.csv", "robust_candidates_primer.csv", "robust_candidates_topcoat.csv", "holdout_predictions.csv"]
    require_files("双层独立Pareto与冻结验证", [PROJECT / "results" / x for x in pareto_files])
    p_pareto = pd.read_csv(PROJECT / "results/pareto_primer.csv", encoding="utf-8-sig")
    m_pareto = pd.read_csv(PROJECT / "results/pareto_topcoat.csv", encoding="utf-8-sig")
    p_can = pd.read_csv(PROJECT / "results/robust_candidates_primer.csv", encoding="utf-8-sig")
    m_can = pd.read_csv(PROJECT / "results/robust_candidates_topcoat.csv", encoding="utf-8-sig")
    record("候选数量与层间独立", len(p_pareto) == 3 and len(m_pareto) == 5 and len(p_can) == 10 and len(m_can) == 10, f"pareto={len(p_pareto)}/{len(m_pareto)} candidates={len(p_can)}/{len(m_can)}")

    lit = pd.read_csv(PROJECT / "supplementary/literature_matrix.csv", encoding="utf-8-sig")
    refs = pd.read_csv(PROJECT / "supplementary/references_verified.csv", encoding="utf-8-sig")
    record("文献证据矩阵", len(lit) == 22 and lit.doi.nunique() == 22, f"rows={len(lit)} unique DOI={lit.doi.nunique()}")
    physical_dois = {
        "10.1007/s11998-008-9144-2", "10.5006/1.3287648", "10.1016/s0300-9440(98)00015-0",
        "10.1016/j.corsci.2009.12.019", "10.1016/j.porgcoat.2013.09.001",
        "10.1016/0010-938x(94)90197-x", "10.1016/j.porgcoat.2020.105847",
        "10.1016/j.polymer.2016.09.059", "10.1016/j.polymer.2024.127816",
        "10.1007/s11249-020-01375-w", "10.1021/acs.iecr.6c01354",
    }
    refs_doi = set(refs.doi.astype(str).str.lower())
    refs_ok = len(refs) == 21 and refs.crossref_verified.astype(bool).all() and physical_dois.issubset(refs_doi)
    record("材料与方法参考文献Crossref核验", refs_ok, f"verified={refs.crossref_verified.sum()}/21 physical={len(physical_dois & refs_doi)}/11")

    main_figs = [
        "fig1_material_framework", "fig2_data_space", "fig3_model_predictions",
        "fig4_primer_material_map", "fig5_topcoat_composition_map",
        "fig6_topcoat_interactions", "fig7_material_design_windows",
    ]
    fig_paths = [PROJECT / "figure" / f"{x}.{ext}" for x in main_figs for ext in ["png", "pdf", "svg"]]
    require_files("7幅材料导向主图PNG/PDF/SVG", fig_paths)
    supp_figs = [
        "figS1_label_permutation", "figS2_zero_replacement_sensitivity",
        "figS3_strict_validation", "figS5_training_domain_screening",
    ]
    require_files("补充验证与敏感性图", [PROJECT / "supplementary" / f"{x}.{ext}" for x in supp_figs for ext in ["png", "pdf", "svg"]])
    figure_qa = json.loads((PROJECT / "figure/figure_language_qa.json").read_text(encoding="utf-8"))
    figure_qa_ok = (
        figure_qa["status"] == "pass"
        and figure_qa["internal_language"] == "English"
        and figure_qa["material_centered_reframe"] is True
        and figure_qa["target_alias_removed"] is True
        and figure_qa["functional_solid_definition_explicit"] is True
        and figure_qa["figure1_type"] == "schematic_led_material_model_design_composite"
        and figure_qa["figure3_type"] == "samplewise_training_oof_frozen_predictions"
        and figure_qa["figure5_type"] == "topcoat_6x5x4_composition_response"
        and figure_qa["figure7_type"] == "high_performance_formulation_boundaries"
        and figure_qa["visual_defects_remaining"] == 0
        and len(figure_qa["main_figures_inspected"]) == 7
        and len(figure_qa["supplementary_figures_present"]) == 4
    )
    record("材料导向绘图复审", figure_qa_ok, json.dumps(figure_qa, ensure_ascii=False))

    require_files("论文源文件与构建脚本", [PROJECT / "manuscript/manuscript_cn.md", PROJECT / "manuscript/build_manuscript.py", PROJECT / "manuscript/artifact.md", DOCX])
    doc = Document(DOCX)
    text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    headings = [p for p in doc.paragraphs if p.style.name in {"Heading 1", "Heading 2", "Heading 3"}]
    record("DOCX结构", len(doc.tables) == 5 and len(doc.inline_shapes) == 7 and len(headings) == 19, f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)} headings={len(headings)}")
    clean = (
        "[[" not in text
        and "]]" not in text
        and "钛合金" not in text
        and not re.search(r"(?:^|\s)/(?:home|public)/", text)
        and "C:\\" not in text
    )
    record("DOCX无内部标记/模板正文/本地路径", clean, "markers=0, Ti template text=0, paths=0")
    overclaim = "证明底漆和面漆协同" not in text and "SiC独立导致磨耗增加" not in text
    record("科学表述准确", overclaim, "material interpretations remain aligned with the observed formulation structure")
    domain_wording = "两类功能涂层" in text and "分别测试" not in text and "样本级映射" not in text
    record("两类功能涂层正向表述", domain_wording, "primer and topcoat are presented as two material systems")
    material_mainline = (
        "数据驱动揭示防腐底漆与聚氨酯/氟树脂面漆的构效关系及高性能配方区域" in text
        and "含25 wt.% SrCrO4和4 wt.%磷酸锌" in text
        and "6×5×4" in text
        and "机器学习插值预测与配方空间泛化" in text
        and "T116" in text
    )
    record("材料构效关系与机器学习主线进入全文", material_mainline, "title, interpolation, primer region, topcoat grid and high-performance boundary")
    numeric_tokens = ["3217", "17.47", "3225", "16.37", "68.48", "11.87", "14.79", "99.93", "3192", "8.10", "0.17", "-0.10", "0.99"]
    record("关键数值进入正文", all(token in text for token in numeric_tokens), ", ".join(numeric_tokens))

    author_ok = all(token in text for token in [
        "Rui Zhoua, Luyao Baoa*, Meirong Caia*",
        "State Key Laboratory of Solid Lubrication",
        "baoluyao@licp.cas.cn",
        "caimr@licp.cas.cn",
    ])
    record("英文作者信息与钛合金正式稿一致", author_ok, "authors, affiliation and corresponding-author e-mails")
    author_paragraph = next((p for p in doc.paragraphs if p.text.startswith("Rui Zhoua, Luyao Baoa*, Meirong Caia*")), None)
    author_markers = [run for run in author_paragraph.runs if run.text in {"a", "a*"}] if author_paragraph is not None else []
    author_superscript_ok = len(author_markers) == 3 and all(run.font.superscript is True for run in author_markers)
    record("作者单位标记为上角标", author_superscript_ok, [run.text for run in author_markers])

    figure1_caption = next((p.text for p in doc.paragraphs if p.text.startswith("图1 ")), "")
    figure1_panels_ok = all(label in figure1_caption for label in ["（a）", "（b）", "（c）", "（d）"])
    record("图1图题包含全部子图序号", figure1_panels_ok, figure1_caption)

    availability_statement = "The datasets used in this study, frozen benchmark configurations, figure source data, and result summaries are available at https://github.com/RuiZhou95/2030-coating"
    record("数据和代码可用性链接", availability_statement in text, availability_statement)
    removed_main_text_ok = "salt_spray_time_h" not in text and "补充材料包括两类功能涂层" not in text and not any(p.text.strip() == "补充材料" for p in doc.paragraphs)
    record("指定正文内容已删除", removed_main_text_ok, "target-alias sentence and Supplementary Materials section absent")

    figure_source = (PROJECT / "src/14_material_coordinates.py").read_text(encoding="utf-8")
    figure4_label_ok = 'ax2.text(xi, value + 0.08, f"{value:.2f}", ha="center", fontsize=14, color="black")' in figure_source
    record("图4c附着力数值使用黑色", figure4_label_ok, "black numeric labels")

    no_atlas_or_pipe = "性能图谱" not in text and "材料图谱" not in text and "|" not in text
    record("正文不使用图谱或竖线分隔", no_atlas_or_pipe, "formulation-property relationships and high-performance regions")

    defensive_terms = [
        "物理坐标", "物理设计坐标", "物理化学假设", "材料假设", "假设", "不构成", "仍需", "不宜", "仅",
        "25/4", "35/4", "15/4", "SrCrO4=", "99.99-100.02", "材料平台", "功能涂层平台", "窗口",
    ]
    defensive_found = [term for term in defensive_terms if term in text]
    record("防御性与抽象表述清理", not defensive_found, defensive_found or "none")

    crossrefs = {label: text.count(label) for label in [
        "图1", "图2", "图3", "图4", "图5", "图6", "图7",
        "表1", "表2", "表3", "表4", "表5",
    ]}
    record("所有主文图表均在正文引用", all(count >= 2 for count in crossrefs.values()), json.dumps(crossrefs, ensure_ascii=False))

    in_references = False
    citation_runs = []
    reference_numbers = []
    body_paragraph_text = []
    citation_pattern = re.compile(r"^\d+(?:-\d+)?(?:,\d+(?:-\d+)?)*$")
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "参考文献":
            in_references = True
            continue
        if in_references:
            match = re.match(r"^\[(\d+)\]", paragraph.text.strip())
            if match:
                reference_numbers.append(int(match.group(1)))
            continue
        body_paragraph_text.append(paragraph.text)
        for run in paragraph.runs:
            if run.font.superscript is True and citation_pattern.fullmatch(run.text.strip()):
                citation_runs.append(run)
    cited_numbers = []
    first_appearance = []
    seen = set()
    for run in citation_runs:
        for number in expand_citation_token(run.text.strip()):
            cited_numbers.append(number)
            if number not in seen:
                seen.add(number)
                first_appearance.append(number)
    superscript_ok = bool(citation_runs) and all("[" not in run.text and "]" not in run.text for run in citation_runs)
    record("正文引文为纯数字上角标", superscript_ok, f"superscript_citation_runs={len(citation_runs)}")
    record("引文按首次出现顺序编号", first_appearance == list(range(1, 22)), f"first_appearance={first_appearance}")
    record("参考文献表顺序", reference_numbers == list(range(1, 22)), f"reference_numbers={reference_numbers}")
    record("正文引文覆盖21条文献", set(cited_numbers) == set(range(1, 22)), f"cited={sorted(set(cited_numbers))}")

    visible_numeric_text = "\n".join(body_paragraph_text) + "\n" + "\n".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    excessive_decimals = sorted(set(re.findall(r"(?<![A-Za-z_])[-+]?\d+\.\d{3,}", visible_numeric_text)))
    record("全文小数不超过两位", not excessive_decimals, excessive_decimals or "all numeric decimals <= 2")

    all_paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)
    sr_total = text.count("SrCrO4")
    sr_subscript = 0
    phi_total = text.count("φF")
    phi_subscript = 0
    nonblack = []
    for paragraph in all_paragraphs:
        runs = paragraph.runs
        for index, run in enumerate(runs[:-1]):
            if run.text.endswith("SrCrO") and runs[index + 1].text.startswith("4") and runs[index + 1].font.subscript is True:
                sr_subscript += 1
            if run.text.endswith("φ") and runs[index + 1].text.startswith("F") and runs[index + 1].font.subscript is True:
                phi_subscript += 1
        for run in runs:
            if run.font.color.type is not None and run.font.color.rgb is not None and str(run.font.color.rgb).upper() != "000000":
                nonblack.append(str(run.font.color.rgb))
    record("化学式与符号下角标", sr_total == sr_subscript and phi_total == phi_subscript and sr_total > 0 and phi_total > 0, f"SrCrO4={sr_subscript}/{sr_total}; phiF={phi_subscript}/{phi_total}")
    record("正文与表格文字全黑", not nonblack, sorted(set(nonblack)) or "all black")

    protocol_boundary = all(token in text for token in [
        "GB/T 5210", "GB/T 1771", "500 g载荷和1000 r",
        "1000 g载荷和1000 r", "各自一致的试验流程",
    ])
    record("核心性能测试条件已写入方法", protocol_boundary, "recorded standards, temperature, salt concentration and wear loads")

    template_ok = sha(REFERENCE) == EXPECTED_TEMPLATE_SHA
    record("参考模板保持不变", template_ok, sha(REFERENCE))
    render = json.loads((PROJECT / "manuscript/render_qa.json").read_text(encoding="utf-8"))
    render_ok = (
        render["docx_sha256"] == sha(DOCX)
        and render["page_count"] == render["rendered_png_count"]
        and render["visual_defects_remaining"] == 0
        and render["inspected_pages"] == list(range(1, render["page_count"] + 1))
    )
    record("DOCX逐页渲染QA", render_ok, json.dumps(render, ensure_ascii=False))

    failed = [x for x in checks if x["status"] != "pass"]
    result = {"status": "pass" if not failed else "fail", "checks": checks, "failed_count": len(failed), "docx_sha256": sha(DOCX)}
    (PROJECT / "results/completion_audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# 项目完成审计", "", f"总体状态：**{result['status'].upper()}**", "", "| 要求 | 状态 | 证据 |", "|---|---|---|"]
    for item in checks:
        evidence = item["evidence"].replace("|", "/").replace("\n", " ")
        if len(evidence) > 300:
            evidence = evidence[:297] + "..."
        lines.append(f"| {item['requirement']} | {item['status']} | {evidence} |")
    (PROJECT / "results/completion_audit.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

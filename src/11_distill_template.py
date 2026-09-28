#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Distill page and style evidence from the Ti manuscript reference."""

from collections import Counter
from hashlib import sha256

from docx import Document
from docx.shared import Emu

from project_paths import PROJECT_ROOT as PROJECT, TEMPLATE_DOCX as REFERENCE

OUT = PROJECT / "manuscript/artifact.md"


def digest(path):
    h = sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def cm(value):
    return round(float(Emu(value).cm), 3) if value is not None else None


def font_value(style, attr):
    value = getattr(style.font, attr)
    if attr == "size" and value is not None:
        return round(value.pt, 2)
    return value


def main():
    doc = Document(REFERENCE)
    style_counts = Counter(p.style.name for p in doc.paragraphs if p.text.strip())
    styles = []
    wanted = ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption", "Abstract", "Keywords"]
    for name in wanted:
        if name not in [s.name for s in doc.styles]:
            continue
        s = doc.styles[name]
        pf = s.paragraph_format
        styles.append(
            f"- {name}: font={font_value(s,'name')}, size={font_value(s,'size')} pt, "
            f"bold={font_value(s,'bold')}, italic={font_value(s,'italic')}, "
            f"alignment={pf.alignment}, before={getattr(pf.space_before,'pt',None)}, "
            f"after={getattr(pf.space_after,'pt',None)}, line_spacing={pf.line_spacing}"
        )
    sections = []
    for i, s in enumerate(doc.sections, 1):
        sections.append(
            f"- Section {i}: page={cm(s.page_width)}×{cm(s.page_height)} cm, "
            f"margins(top/bottom/left/right)={cm(s.top_margin)}/{cm(s.bottom_margin)}/{cm(s.left_margin)}/{cm(s.right_margin)} cm, "
            f"header={cm(s.header_distance)} cm, footer={cm(s.footer_distance)} cm"
        )
    headers = [" / ".join(p.text.strip() for p in s.header.paragraphs if p.text.strip()) for s in doc.sections]
    footers = [" / ".join(p.text.strip() for p in s.footer.paragraphs if p.text.strip()) for s in doc.sections]
    lines = [
        "# 模板提取执行契约",
        "",
        "## Reference",
        "",
        f"- 远程参考：`{REFERENCE}`",
        f"- SHA256：`{digest(REFERENCE)}`",
        f"- 大小：{REFERENCE.stat().st_size} bytes",
        "- Word COM 渲染页数：16页（只读渲染证据保存在本次任务临时QA目录）",
        f"- Section数：{len(doc.sections)}",
        f"- 表格数：{len(doc.tables)}；内嵌图数：{len(doc.inline_shapes)}",
        "",
        "## Page system",
        "",
        *sections,
        "",
        "## Typography evidence",
        "",
        *styles,
        "",
        "### 正文中实际使用的样式计数",
        "",
        *[f"- {k}: {v}" for k, v in style_counts.most_common()],
        "",
        "## Header/footer",
        "",
        f"- Header text: {headers}",
        f"- Footer text: {footers}",
        "",
        "## Content flow and slot map",
        "",
        "参考稿呈现：中文题名与作者占位、中文摘要/关键词、英文摘要/关键词、引言、数据与方法、结果与讨论、结论、声明与数字参考文献。",
        "本任务不是窄幅改写：用户明确要求只参考版芯、字体、标题层级、图表题注和参考文献格式，并禁止复制钛合金论文叙事。因此新稿将保留页面系统和可复用样式，完整替换正文内容、表格和图像。",
        "",
        "## Package preservation",
        "",
        "- 参考DOCX保持只读且不被修改。",
        "- 新稿从参考文件的工作副本构建，保留styles、numbering、section、theme、header/footer关系；正文内容、表格、图片关系按新论文重建。",
        "- 由于正文长度和图表数量不同，分页与参考稿不同属于预期变化，不作为模板保真失败。",
        "",
        "## Fidelity gates",
        "",
        "1. 参考文件SHA256在交付前复核不变。",
        "2. 新稿保持相同A4页面与页边距、正文和标题字体体系、题注和数字参考文献风格。",
        "3. 最终稿通过DOCX结构审计，并使用Word COM导出PDF后逐页检查。",
        "4. 不保留任何钛合金正文、表格或图像内容。",
    ]
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"TEMPLATE_DISTILLED sha256={digest(REFERENCE)} sections={len(doc.sections)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()

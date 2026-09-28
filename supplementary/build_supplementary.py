#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the standalone Chinese Supporting Information DOCX."""

from __future__ import annotations

import os
import re
from decimal import Decimal
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT = Path(
    os.environ.get("COATING_PROJECT_ROOT", Path(__file__).resolve().parents[1])
).expanduser().resolve()
SOURCE = PROJECT / "supplementary/supplementary_cn.md"
OUTPUT = PROJECT / "supplementary/Supplementary_Information_Coating_CN.docx"
DATA = PROJECT / "data"
RESULTS = PROJECT / "results"
SUPP = PROJECT / "supplementary"

FIGURES = {
    "[[FIG_S1]]": (
        SUPP / "figS5_training_domain_screening.png",
        "图S1 两类功能涂层配方的实验复核优先级。排位综合实测性能、偏向不利方向的性能保守估计和设计域距离；排位越靠前，越适合优先重复制备与性能复测。",
    ),
    "[[FIG_S2]]": (
        SUPP / "figS3_strict_validation.png",
        "图S2 最终模型在重复随机划分、配方类别整体留出和关键组成水平整体留出下的外层测试R²。底漆重复随机划分体现已覆盖组成之间的插值性能，面漆盐雾和磨耗在三类划分下均保持稳定。",
    ),
    "[[FIG_S3]]": (
        SUPP / "figS1_label_permutation.png",
        "图S3 标签置换负对照。箱线图为100次标签置换在结构化留出条件下的R²分布，红点为原始目标的结构化留出R²。",
    ),
    "[[FIG_S4]]": (
        SUPP / "figS2_zero_replacement_sensitivity.png",
        "图S4 面漆ILR表示对零替代因子的敏感性。零组分分别以最小正组分的0.10、0.50和0.90倍替代并重新闭合，曲线比较配方类别整体留出和关键组成水平整体留出下的R²变化。",
    ),
}

SCIENTIFIC_RE = re.compile(r"(SrCrO4|φF)")


def normalize(text: str) -> str:
    return text.replace("–", "-").replace("‑", "-").replace("−", "-").replace("`", "")


def set_run_font(run, east="宋体", latin="Times New Roman", size=10.5, bold=None, italic=None, color=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:ascii"), latin)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), latin)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(*(color or (0, 0, 0)))


def add_scientific_text(paragraph, text, size=10.5, east="宋体", latin="Times New Roman", bold=None, italic=None):
    for token in SCIENTIFIC_RE.split(normalize(text)):
        if not token:
            continue
        if token == "SrCrO4":
            run = paragraph.add_run("SrCrO")
            set_run_font(run, east=east, latin=latin, size=size, bold=bold, italic=italic)
            sub = paragraph.add_run("4")
            set_run_font(sub, east=east, latin=latin, size=size, bold=bold, italic=italic)
            sub.font.subscript = True
        elif token == "φF":
            run = paragraph.add_run("φ")
            set_run_font(run, east=east, latin=latin, size=size, bold=bold, italic=italic)
            sub = paragraph.add_run("F")
            set_run_font(sub, east=east, latin=latin, size=size, bold=bold, italic=italic)
            sub.font.subscript = True
        else:
            run = paragraph.add_run(token)
            set_run_font(run, east=east, latin=latin, size=size, bold=bold, italic=italic)


def set_cell_text(cell, value):
    paragraph = cell.paragraphs[0]
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    add_scientific_text(paragraph, str(value), size=8.0)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=80, bottom=70, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_widths(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "80")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        node = OxmlElement("w:gridCol")
        node.set(qn("w:w"), str(int(width * 567)))
        grid.append(node)
    for row in table.rows:
        prevent_split(row)
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def format_table(table, widths, font_size=7.6, alignments=None):
    table.style = "Table Grid"
    repeat_header(table.rows[0])
    set_table_widths(table, widths)
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if i == 0:
                shade_cell(cell, "D9EAF7")
            elif i % 2 == 0:
                shade_cell(cell, "F7FAFC")
            for p in cell.paragraphs:
                p.alignment = (alignments[j] if alignments else WD_ALIGN_PARAGRAPH.CENTER)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=(i == 0))


def add_table_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    set_run_font(p.add_run(text), east="黑体", size=9.2, bold=True)


def make_table(doc, title, rows, widths, font_size=7.6, alignments=None):
    add_table_title(doc, title)
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            set_cell_text(table.cell(i, j), value)
    format_table(table, widths, font_size=font_size, alignments=alignments)
    return table


def table_s1(doc):
    stats = pd.read_csv(DATA / "target_statistics.csv", encoding="utf-8-sig")
    rows = [["材料体系", "样本/结构", "响应", "n/唯一值", "范围", "数据角色"]]
    mapping = {
        ("底漆", "salt_spray_pass_h"): ("盐雾失效时间/h", "核心"),
        ("底漆", "adhesion_mpa_or_grade"): ("拉开附着力/MPa", "核心"),
        ("底漆", "wear_mass_loss_mg"): ("磨耗失重/mg", "辅助"),
        ("面漆", "salt_spray_pass_h"): ("盐雾失效时间/h", "核心"),
        ("面漆", "wear_mass_loss_mg"): ("磨耗失重/mg", "核心"),
    }
    for _, r in stats.iterrows():
        key = (r["layer"], r["target"])
        if key not in mapping:
            continue
        response, role = mapping[key]
        structure = "72；名义4×3×3×2，71唯一组合" if r["layer"] == "底漆" else "120；完整6×5×4网格"
        rows.append([r["layer"], structure, response, f"{int(r['n'])}/{int(r['unique_n'])}", f"{r['min']:.2f}–{r['max']:.2f}", role])
    make_table(doc, "表S1 两类功能涂层的数据结构与目标范围", rows, [1.5, 3.7, 2.8, 1.9, 2.5, 2.2], 7.4)


def table_s2(doc):
    nested = pd.read_csv(RESULTS / "nested_cv_summary.csv", encoding="utf-8-sig")
    selected = pd.read_csv(RESULTS / "final_model_selection.csv", encoding="utf-8-sig")
    selected = selected[["layer", "target", "target_role", "representation", "model"]]
    q = nested.merge(selected, on=["layer", "target", "representation", "model"])
    target_names = {
        ("primer", "salt_spray_pass_h"): "底漆盐雾",
        ("primer", "adhesion_mpa_or_grade"): "底漆附着",
        ("primer", "wear_mass_loss_mg"): "底漆磨耗(辅助)",
        ("topcoat", "salt_spray_pass_h"): "面漆盐雾",
        ("topcoat", "wear_mass_loss_mg"): "面漆磨耗",
    }
    scheme_names = {"random_repeated": "重复随机", "family_logo": "配方族留出", "level_logo": "关键水平留出"}
    order = {"random_repeated": 0, "family_logo": 1, "level_logo": 2}
    rows = [["响应", "表示/模型", "重复随机\nR²/NRMSE/ρ", "配方类别整体留出\nR²/NRMSE/ρ", "关键组成水平整体留出\nR²/NRMSE/ρ"]]
    for _, s in selected.iterrows():
        g = q[(q["layer"] == s["layer"]) & (q["target"] == s["target"])].set_index("scheme")
        metrics = []
        for scheme in ("random_repeated", "family_logo", "level_logo"):
            r = g.loc[scheme]
            metrics.append(f"{r['pooled_r2']:.2f}/{r['pooled_nrmse_range']:.2f}/{r['pooled_spearman']:.2f}")
        rows.append([
            target_names[(s["layer"], s["target"])],
            f"{s['representation']}/{s['model']}",
            *metrics,
        ])
    make_table(doc, "表S6 最终模型在三类外层验证方案下的汇总性能（R²/NRMSE/Spearman ρ）", rows, [2.5, 2.5, 3.2, 3.2, 3.2], 7.4)


def table_s3(doc):
    rows = [
        ["审计项", "观测状态", "解释"],
        ["名义因子结构", "4×3×3×2=72", "SrCrO4×磷酸锌×固化剂×固化制度"],
        ["记录数/唯一组合", "72/71", "存在1个重复组合和1个对应未覆盖组合"],
        ["重复组合", "SrCrO4 15 wt.%；磷酸锌4 wt.%；酚醛胺；25 ℃/168 h；n=2", "同一名义配方与工艺条件的两条记录"],
        ["未覆盖组合", "SrCrO4 15 wt.%；磷酸锌4 wt.%；酚醛胺；80 ℃/4 h；n=0", "对应组成单元的另一固化制度"],
        ["重点组成单元", "SrCrO4为25或35 wt.%；磷酸锌4 wt.%", "每个单元均覆盖3类固化剂×2种固化制度"],
    ]
    make_table(doc, "表S2 底漆名义因子网格的重复与缺失审计", rows, [3.0, 5.0, 6.6], 7.7, [WD_ALIGN_PARAGRAPH.LEFT]*3)


def table_s4(doc):
    df = pd.read_csv(RESULTS / "zero_replacement_sensitivity.csv", encoding="utf-8-sig")
    rows = [["响应", "零替代因子", "配方类别整体留出R²/ρ", "关键组成水平整体留出R²/ρ"]]
    names = {"salt_spray_pass_h": "面漆盐雾", "wear_mass_loss_mg": "面漆磨耗"}
    for target in ["salt_spray_pass_h", "wear_mass_loss_mg"]:
        for factor in [0.1, 0.5, 0.9]:
            q = df[(df["target"] == target) & (df["zero_factor"] == factor)].set_index("scheme")
            f = q.loc["family_logo"]
            l = q.loc["level_logo"]
            rows.append([names[target], f"{factor:.2f}", f"{f['r2']:.2f}/{f['spearman']:.2f}", f"{l['r2']:.2f}/{l['spearman']:.2f}"])
    make_table(doc, "表S7 面漆ILR零替代敏感性（R²/Spearman ρ）", rows, [2.6, 2.2, 4.3, 4.3], 7.8)


def table_s5(doc):
    df = pd.read_csv(RESULTS / "uncertainty_summary.csv", encoding="utf-8-sig")
    df = df[df["set"] == "frozen_holdout"].copy()
    target_names = {
        ("primer", "adhesion_mpa_or_grade"): "底漆附着",
        ("primer", "salt_spray_pass_h"): "底漆盐雾",
        ("primer", "wear_mass_loss_mg"): "底漆磨耗(辅助)",
        ("topcoat", "salt_spray_pass_h"): "面漆盐雾",
        ("topcoat", "wear_mass_loss_mg"): "面漆磨耗",
    }
    rows = [["响应", "n", "Bootstrap 90%覆盖", "Conformal 90%覆盖", "平均Bootstrap宽度", "Conformal宽度"]]
    for _, r in df.iterrows():
        rows.append([
            target_names[(r["layer"], r["target"])],
            f"{int(r['n'])}",
            f"{r['bootstrap_90_coverage']:.2f}",
            f"{r['conformal_90_coverage']:.2f}",
            f"{r['mean_bootstrap_width']:.2f}",
            f"{r['conformal_width']:.2f}",
        ])
    make_table(doc, "表S6 预先冻结样本的预测区间覆盖", rows, [2.4, 1.2, 2.6, 2.6, 2.9, 2.9], 7.6)


def table_s6(doc):
    df = pd.read_csv(RESULTS / "topcoat_variance_decomposition.csv", encoding="utf-8-sig")
    target_names = {"salt_spray_pass_h": "面漆盐雾", "wear_mass_loss_mg": "面漆磨耗"}
    term_names = {
        "Binder fluorination (phi_F)": "树脂相内氟树脂质量分数φF",
        "Functional-solid/binder ratio (psi)": "功能固相/树脂比ψ",
        "Filler package": "耦合填料包",
        "phi_F x psi": "φF×ψ",
        "phi_F x package": "φF×填料包",
        "psi x package": "ψ×填料包",
        "Three-way interaction": "三阶交互",
    }
    rows = [["响应", "项", "观测平方和", "占总平方和/%"]]
    for _, r in df.iterrows():
        value = float(r["sum_of_squares"])
        if abs(value) >= 1:
            four_sig = f"{value:.4g}"
            sum_of_squares = format(Decimal(four_sig), "f") if "e" in four_sig.lower() else four_sig
        else:
            sum_of_squares = f"{value:.2f}"
        rows.append([target_names[r["target"]], term_names[r["term"]], sum_of_squares, f"{r['variance_pct']:.2f}"])
    make_table(doc, "表S3 面漆完整6×5×4网格的描述性平方和分解", rows, [2.5, 5.0, 3.2, 3.2], 7.5)


def table_s7(doc):
    df = pd.read_csv(RESULTS / "robust_candidates_primer.csv", encoding="utf-8-sig").head(10)
    rows = [["样本", "缓蚀颜料含量", "固化剂", "实测盐雾寿命/附着力", "盐雾寿命/附着力保守估计", "域距离"]]
    for _, r in df.iterrows():
        sid = f"P{int(str(r['sample_id']).split('_')[-1]):02d}"
        rows.append([
            sid,
            f"SrCrO4 {r['铬酸锶']:.0f} wt.%\n磷酸锌 {r['磷酸锌']:.0f} wt.%",
            r["coating_curing_agent_type"],
            f"{r['salt_spray_pass_h']:.0f} h/{r['adhesion_mpa_or_grade']:.2f} MPa",
            f"{r['salt_spray_pass_h__robust']:.0f} h/{r['adhesion_mpa_or_grade__robust']:.2f} MPa",
            f"{r['domain_distance']:.2f}",
        ])
    make_table(doc, "表S4 底漆训练域按实验复核优先级排列的10个既有样本", rows, [1.5, 3.1, 1.7, 2.8, 2.8, 1.7], 6.6)


def table_s8(doc):
    df = pd.read_csv(RESULTS / "robust_candidates_topcoat.csv", encoding="utf-8-sig").head(10)
    rows = [["样本", "树脂含量", "功能固相含量", "实测盐雾寿命/磨耗失重", "盐雾寿命/磨耗失重保守估计", "域距离"]]
    for _, r in df.iterrows():
        sid = f"T{int(str(r['sample_id']).replace('面漆', '')):03d}"
        rows.append([
            sid,
            f"PU {r['聚氨酯树脂A14含量']:.2f} wt.%\nF {r['氟树脂含量']:.2f} wt.%",
            f"SiC {r['碳化硅']:.2f}\nBN {r['氮化硼']:.2f}\n润滑粉 {r['润滑粉']:.2f} wt.%",
            f"{r['salt_spray_pass_h']:.0f} h/{r['wear_mass_loss_mg']:.2f} mg",
            f"{r['salt_spray_pass_h__robust']:.0f} h/{r['wear_mass_loss_mg__robust']:.2f} mg",
            f"{r['domain_distance']:.2f}",
        ])
    make_table(doc, "表S5 面漆训练域按实验复核优先级排列的10个既有样本", rows, [1.5, 2.7, 3.0, 2.4, 2.7, 1.7], 6.2)


TABLES = {
    "[[TABLE_S1]]": table_s1,
    "[[TABLE_S2]]": table_s3,
    "[[TABLE_S3]]": table_s6,
    "[[TABLE_S4]]": table_s7,
    "[[TABLE_S5]]": table_s8,
    "[[TABLE_S6]]": table_s5,
    "[[TABLE_S7]]": table_s4,
}


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    add_scientific_text(p, text, size=8.8, bold=True)


def add_figure(doc, path, caption):
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Cm(14.2))
    shape._inline.docPr.set("descr", normalize(caption))
    add_caption(doc, caption)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("S")
    set_run_font(run, size=8.5)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_document(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.175)
    sec.header_distance = sec.footer_distance = Cm(1.27)
    add_page_number(sec.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(3)

    for name, size, before, after in (("Heading 1", 14, 14, 5), ("Heading 2", 11.5, 10, 4), ("Heading 3", 10.5, 8, 3)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    add_scientific_text(p, text, size=10.5)


def build():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "防腐涂层构效关系与高性能配方区域支撑信息"
    doc.core_properties.subject = "底漆与面漆构效关系、候选配方和模型验证"
    doc.core_properties.keywords = "支撑信息; 防腐底漆; 聚氨酯/氟树脂面漆; 构效关系; 机器学习; 配方区域"

    first_title = True
    for raw in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line in TABLES:
            TABLES[line](doc)
            continue
        if line in FIGURES:
            add_figure(doc, *FIGURES[line])
            continue
        if line == "[[PAGEBREAK]]":
            doc.add_page_break()
            continue
        if line.startswith("# "):
            text = normalize(line[2:])
            if first_title:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(7)
                p.paragraph_format.keep_with_next = True
                set_run_font(p.add_run(text), east="黑体", size=17, bold=True)
                first_title = False
            else:
                doc.add_paragraph(text, style="Heading 1")
            continue
        if line.startswith("## "):
            text = normalize(line[3:])
            if text.startswith("Data-Driven Formulation-Property Relationships"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.keep_with_next = True
                set_run_font(p.add_run(text), east="黑体", size=12.5, bold=True)
            else:
                doc.add_paragraph(text, style="Heading 2")
            continue
        if line.startswith("### "):
            text = normalize(line[4:])
            if text.startswith("Supporting Information"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                set_run_font(p.add_run(text), size=10.5, bold=True)
            else:
                doc.add_paragraph(text, style="Heading 3")
            continue
        if line.startswith("Rui Zhou"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            authors = [("Rui Zhou", "a"), ("Luyao Bao", "a*"), ("Meirong Cai", "a*")]
            for index, (name, marker) in enumerate(authors):
                set_run_font(p.add_run(name), size=9.5)
                marker_run = p.add_run(marker)
                set_run_font(marker_run, size=9.5)
                marker_run.font.superscript = True
                if index < len(authors) - 1:
                    set_run_font(p.add_run(", "), size=9.5)
            continue
        if line.startswith(("a. State Key Laboratory", "*Corresponding author")):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            add_scientific_text(p, line, size=9.5, italic=line.startswith("*Corresponding"))
            continue
        if line.endswith("  ") or (line.startswith("补充") and "　" in line):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(2)
            add_scientific_text(p, line.rstrip(), size=9.5)
            continue
        add_body(doc, line)

    doc.save(OUTPUT)
    print(
        f"SI_BUILT path={OUTPUT} bytes={OUTPUT.stat().st_size} "
        f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}"
    )


if __name__ == "__main__":
    build()

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the Chinese SCI-style coating manuscript as a clean DOCX."""

from __future__ import annotations

import html
import json
import os
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PROJECT = Path(
    os.environ.get("COATING_PROJECT_ROOT", Path(__file__).resolve().parents[1])
).expanduser().resolve()
SOURCE = PROJECT / "manuscript/manuscript_cn.md"
OUTPUT = PROJECT / "manuscript/Manuscript_Coating_CN.docx"
RESULTS = PROJECT / "results"
FIG = PROJECT / "figure"
REFS = PROJECT / "supplementary/references_verified.csv"

FIGURES = {
    "[[FIGURE1]]": (FIG / "fig1_material_framework.png", "图1 涂层组成、材料结构、机器学习映射与配方设计的研究框架。（a）环氧防腐底漆的缓蚀颜料、环氧基体与转化层界面；（b）聚氨酯/氟树脂面漆的树脂连续相、功能固相与固定底层；（c）关键材料变量、性能响应、组成表示和嵌套模型选择；（d）响应面、高性能边界与代表性配方输出。", 14.5),
    "[[FIGURE2]]": (FIG / "fig2_data_space.png", "图2 两类功能涂层的实测性能与原始配方空间。（a,b）底漆盐雾寿命与附着力分布；（c）SrCrO4—磷酸锌设计网格；（d,e）面漆盐雾寿命与磨耗分布；（f）SiC—BN组分空间，标记面积表示氟树脂含量。"),
    "[[FIGURE3]]": (FIG / "fig3_model_predictions.png", "图3 训练样本、配方类别整体留出的折外测试样本和预先冻结样本的逐样本预测。（a）底漆盐雾寿命；（b）底漆拉开附着力；（c）面漆盐雾寿命；（d）面漆磨耗失重。灰色圆点表示模型拟合训练样本，蓝色三角表示折外测试预测，红色圆点表示冻结样本。红色误差线为200次Bootstrap预测的5%-95%分位区间；上下分位数分别由预测分布计算，因此区间可相对点预测呈非对称。"),
    "[[FIGURE4]]": (FIG / "fig4_primer_material_map.png", "图4 底漆缓蚀颜料优化区与固化/界面条件。（a,b）SrCrO4—磷酸锌组成单元的盐雾和附着均值；（c）不同固化剂的性能基线；（d）表面粗糙度与盐雾寿命的实测关系。"),
    "[[FIGURE5]]": (FIG / "fig5_topcoat_composition_map.png", "图5 面漆树脂组成与功能固相负载的性能响应。（a）盐雾寿命均值；（b）磨耗失重均值；（c）6×5×4网格观测总平方和的描述性分解。"),
    "[[FIGURE6]]": (FIG / "fig6_topcoat_interactions.png", "图6 面漆关键配方变量的交互与代表性材料。（a）不同φF条件下ψ对盐雾寿命的影响；（b）不同φF条件下ψ对磨耗的影响；（c）T108、T112、T68和T116的相组成及实测性能。"),
    "[[FIGURE7]]": (FIG / "fig7_material_design_windows.png", "图7 两类功能涂层的高性能配方边界。（a）底漆SrCrO4—磷酸锌组成单元及较低SrCrO4负载路线P42；（b）面漆完整配方空间的盐雾—磨耗关系。颜色表示φF，点面积随ψ降低而增大。"),
}

CITATION_RE = re.compile(r"(\[(?:\d+(?:-\d+)?)(?:,\d+(?:-\d+)?)*\])")
SCIENTIFIC_RE = re.compile(r"(SrCrO4|φF)")
REFERENCE_NUMBER_MAP = {
    11: 1, 12: 2, 13: 3, 14: 4, 15: 5, 16: 6,
    18: 7, 19: 8, 20: 9, 21: 10,
    1: 11, 2: 12, 3: 13, 4: 14, 5: 15,
    6: 16, 7: 17, 8: 18, 9: 19, 10: 20, 17: 21,
}


def normalize(text: str) -> str:
    return (
        text.replace("–", "-").replace("‑", "-")
        .replace("−", "-").replace("↑", "增大").replace("↓", "减小")
        .replace("`", "").replace("$", "")
    )


def set_run_font(run, east="宋体", latin="Times New Roman", size=10.5, bold=None, italic=None, color=None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor(*(color or (0, 0, 0)))


def add_scientific_text(paragraph, text, size=10.5, east="宋体", latin="Times New Roman", bold=None, italic=None):
    """Add text with true subscripts for the recurring chemical and symbol notation."""
    for token in SCIENTIFIC_RE.split(normalize(text)):
        if not token:
            continue
        if token == "SrCrO4":
            run = paragraph.add_run("SrCrO")
            set_run_font(run, east=east, latin=latin, size=size, bold=bold, italic=italic)
            sub = paragraph.add_run("4")
            set_run_font(sub, east=east, latin=latin, size=size, bold=bold, italic=italic)
            sub.font.subscript = True
        elif token == "φF":
            run = paragraph.add_run("φ")
            set_run_font(run, east=east, latin=latin, size=size, bold=bold, italic=italic)
            sub = paragraph.add_run("F")
            set_run_font(sub, east=east, latin=latin, size=size, bold=bold, italic=italic)
            sub.font.subscript = True
        else:
            run = paragraph.add_run(token)
            set_run_font(run, east=east, latin=latin, size=size, bold=bold, italic=italic)


def add_text_with_superscript_citations(paragraph, text, size=10.5, east="宋体", latin="Times New Roman", bold=None, italic=None):
    """Add prose while formatting numeric square-bracket citations as superscript runs."""
    for part in CITATION_RE.split(normalize(text)):
        if not part:
            continue
        if CITATION_RE.fullmatch(part):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, east=east, latin=latin, size=size, bold=bold, italic=italic)
            run.font.superscript = True
        else:
            add_scientific_text(paragraph, part, size=size, east=east, latin=latin, bold=bold, italic=italic)


def set_cell_text(cell, value):
    paragraph = cell.paragraphs[0]
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    add_scientific_text(paragraph, str(value), size=8.0)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def prevent_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def set_table_widths(table, widths_cm):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total_twips = int(sum(widths_cm) * 567)
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in("w:tblW")
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total_twips))
    tblW.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(width * 567)))
        grid.append(gc)
    for row in table.rows:
        prevent_split(row)
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.first_child_found_in("w:tcW")
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 567)))
            tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def format_table(table, header_fill="D9EAF7", font_size=8.0):
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if i == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), header_fill)
                cell._tc.get_or_add_tcPr().append(shd)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, size=font_size, bold=(i == 0))


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_with_next = False
    add_scientific_text(p, text, size=9.0, bold=True)
    return p


def add_figure(doc, path, caption, width_cm=14.2):
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Cm(width_cm))
    docPr = shape._inline.docPr
    docPr.set("descr", normalize(caption))
    add_caption(doc, caption)


def add_table_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    set_run_font(p.add_run(text), size=9.0, bold=True)


def table1(doc):
    add_table_title(doc, "表1 两类功能涂层材料体系及其配方结构")
    data = [
        ["材料体系", "基材与背景", "涂层与固化", "配方结构", "性能响应"],
        ["环氧底漆", "2024铝板；打磨后转化处理", "25-30 μm；3类固化剂；2种固化制度", "名义4×3×3×2；72条记录、71个唯一组合", "盐雾失效时间；拉开附着力"],
        ["PU/氟树脂面漆", "钢材；喷砂除锈；固定底层构型", "异氰酸酯固化；八组分配方", "φF 6级×ψ 5级×填料包4级", "盐雾失效时间；磨耗失重"],
    ]
    t = doc.add_table(rows=len(data), cols=len(data[0]))
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            set_cell_text(t.cell(i, j), value)
    set_table_widths(t, [2.0, 3.0, 3.6, 3.5, 2.4])
    format_table(t, font_size=7.8)


def table2(doc):
    add_table_title(doc, "表2 面向涂层构效关系的关键配方变量")
    data = [
        ["体系", "关键变量", "定义或水平", "材料含义"],
        ["底漆", "缓蚀颜料用量", "SrCrO4为0、15、25或35 wt.%；磷酸锌为0、4或8 wt.%", "活性缓蚀物种供给与颜料负载"],
        ["底漆", "固化与界面条件", "3类固化剂；25 ℃/168 h或80 ℃/4 h；Ra=1.5-3.5 μm", "网络化学、颜料润湿和界面结合"],
        ["面漆", "树脂相内氟树脂质量分数φF", "F/(PU+F)=0、0.25、0.33、0.50、0.75或1", "连续树脂相的化学组成"],
        ["面漆", "功能固相/树脂质量比ψ", "(SiC+BN+润滑粉)/(PU+F)=0.25、0.43、1、2.33或4", "树脂连续性与功能固相负载"],
        ["面漆", "耦合填料包", "BN陶瓷占比与润滑粉固相占比联动，共4级", "承载、层间剪切与润滑的组合变化"],
    ]
    t = doc.add_table(rows=len(data), cols=4)
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            set_cell_text(t.cell(i, j), value)
    set_table_widths(t, [1.8, 3.0, 5.2, 4.6])
    format_table(t, font_size=7.7)


def table3(doc):
    add_table_title(doc, "表3 最终模型在三类外层验证方案下的汇总性能（R²/NRMSE/Spearman ρ）")
    nested = pd.read_csv(RESULTS / "nested_cv_summary.csv", encoding="utf-8-sig")
    selected = pd.read_csv(RESULTS / "final_model_selection.csv", encoding="utf-8-sig")
    selected = selected[["layer", "target", "target_role", "representation", "model"]]
    merged = nested.merge(selected, on=["layer", "target", "representation", "model"])
    target_names = {
        ("primer", "salt_spray_pass_h"): "底漆盐雾",
        ("primer", "adhesion_mpa_or_grade"): "底漆附着",
        ("primer", "wear_mass_loss_mg"): "底漆磨耗（辅助）",
        ("topcoat", "salt_spray_pass_h"): "面漆盐雾",
        ("topcoat", "wear_mass_loss_mg"): "面漆磨耗",
    }
    data = [["响应", "表示/模型", "重复随机\nR²/NRMSE/ρ", "配方类别整体留出\nR²/NRMSE/ρ", "关键组成水平整体留出\nR²/NRMSE/ρ"]]
    for _, selection in selected.iterrows():
        group = merged[(merged["layer"] == selection["layer"]) & (merged["target"] == selection["target"])].set_index("scheme")
        metrics = []
        for scheme in ("random_repeated", "family_logo", "level_logo"):
            row = group.loc[scheme]
            metrics.append(f"{row['pooled_r2']:.2f}/{row['pooled_nrmse_range']:.2f}/{row['pooled_spearman']:.2f}")
        data.append([
            target_names[(selection["layer"], selection["target"])],
            f"{selection['representation']}/{selection['model']}",
            *metrics,
        ])
    table = doc.add_table(rows=len(data), cols=5)
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            set_cell_text(table.cell(i, j), value)
    set_table_widths(table, [2.2, 3.0, 3.1, 3.1, 3.1])
    format_table(table, font_size=7.3)


def table4(doc):
    add_table_title(doc, "表4 底漆缓蚀颜料与固化剂的实测边际规律")
    factor = pd.read_csv(RESULTS / "primer_factor_means.csv", encoding="utf-8-sig")
    data = [["因素", "水平", "盐雾均值/h", "附着力均值/MPa", "材料设计解读"]]
    interpretations = {
        ("铬酸锶", "0.0"): "基础颜料条件",
        ("铬酸锶", "15.0"): "盐雾和附着明显提高",
        ("铬酸锶", "25.0"): "盐雾进入稳定区且附着最高",
        ("铬酸锶", "35.0"): "盐雾进入稳定区，附着回落",
        ("磷酸锌", "0.0"): "基础颜料条件",
        ("磷酸锌", "4.0"): "两项性能均出现内部最大值",
        ("磷酸锌", "8.0"): "相对4 wt.%两项性能均回落",
        ("coating_curing_agent_type", "改性胺"): "较高耐蚀—附着性能基线",
        ("coating_curing_agent_type", "聚酰胺"): "中等性能基线",
        ("coating_curing_agent_type", "酚醛胺"): "较低耐蚀—附着性能基线",
    }
    factor_names = {"铬酸锶": "SrCrO4", "磷酸锌": "Zn phosphate", "coating_curing_agent_type": "固化剂"}
    for _, row in factor.iterrows():
        level = str(row["level"])
        key = (row["factor"], level)
        data.append([
            factor_names.get(row["factor"], row["factor"]),
            level,
            f"{row['salt_spray_pass_h']:.0f}",
            f"{row['adhesion_mpa_or_grade']:.2f}",
            interpretations.get(key, "当前数据域的描述性均值"),
        ])
    t = doc.add_table(rows=len(data), cols=5)
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            set_cell_text(t.cell(i, j), value)
    set_table_widths(t, [2.1, 2.0, 2.4, 2.7, 5.4])
    format_table(t, font_size=7.5)


def table5(doc):
    top = pd.read_csv(RESULTS / "topcoat_physical_coordinates.csv", encoding="utf-8-sig")
    add_table_title(doc, "表5 代表性配方及其材料设计定位")
    data = [["体系/材料", "关键配方参数", "实测性能", "材料设计定位"]]
    p42 = pd.read_csv(RESULTS / "robust_candidates_primer.csv", encoding="utf-8-sig")
    p42 = p42[p42["sample_id"].astype(str).str.endswith("42")].iloc[0]
    data.append([
        "底漆P42",
        f"SrCrO4 {p42['铬酸锶']:.0f} wt.%；磷酸锌 {p42['磷酸锌']:.0f} wt.%；{p42['coating_curing_agent_type']}",
        f"3050 h；16.90 MPa",
        "较低SrCrO4负载的优先复核路线",
    ])
    roles = {108: "磨耗优先", 112: "平衡型", 68: "盐雾优先", 116: "完整域性能上界"}
    for suffix in [108, 112, 68, 116]:
        r = top[top["sample_id"].astype(str).str.endswith(str(suffix))].iloc[0]
        data.append([
            f"面漆T{suffix}",
            f"φF={r['fluororesin_fraction_in_binder']:.2f}; ψ={r['solid_to_binder_mass_ratio']:.2f}; 包{int(r['filler_package_level'])}",
            f"{r['salt_spray_pass_h']:.0f} h；{r['wear_mass_loss_mg']:.2f} mg",
            roles[suffix],
        ])
    t = doc.add_table(rows=len(data), cols=4)
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            set_cell_text(t.cell(i, j), value)
    set_table_widths(t, [2.2, 5.3, 3.0, 4.4])
    format_table(t, font_size=7.7)


TABLES = {
    "[[TABLE1]]": table1,
    "[[TABLE2]]": table2,
    "[[TABLE3]]": table3,
    "[[TABLE4]]": table4,
    "[[TABLE5]]": table5,
}


def reference_text(row, display_number):
    authors = html.unescape(str(row.authors or "")).strip()
    title = html.unescape(str(row.title or "")).strip()
    journal = html.unescape(str(row.journal or "")).strip()
    year = str(int(row.year)) if pd.notna(row.year) and str(row.year) != "" else ""
    volume = "" if pd.isna(row.volume) else str(row.volume)
    issue = "" if pd.isna(row.issue) else str(row.issue)
    pages = "" if pd.isna(row.pages_or_article) else str(row.pages_or_article)
    vol = volume + (f"({issue})" if issue else "")
    bibliographic = f"{journal}. {year}"
    if vol:
        bibliographic += f"; {vol}"
    if pages:
        bibliographic += f": {pages}"
    return normalize(f"[{int(display_number)}] {authors}. {title}. {bibliographic}. https://doi.org/{row.doi}")


def add_references(doc):
    refs = pd.read_csv(REFS, encoding="utf-8-sig")
    refs["display_number"] = refs["number"].map(REFERENCE_NUMBER_MAP)
    if refs["display_number"].isna().any() or set(refs["display_number"].astype(int)) != set(range(1, 22)):
        raise AssertionError("Reference renumbering does not form a complete 1-21 sequence")
    refs = refs.sort_values("display_number")
    for _, row in refs.iterrows():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        set_run_font(p.add_run(reference_text(row, row["display_number"])), size=8.5)


def configure_styles(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.175)
    sec.header_distance = sec.footer_distance = Cm(1.27)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.first_line_indent = Cm(0.74)
    pf.line_spacing = 1.15
    pf.space_after = Pt(3)

    for name, size, east in [("Heading 1", 14, "黑体"), ("Heading 2", 11.5, "黑体"), ("Heading 3", 10.5, "黑体")]:
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), east)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.first_line_indent = Cm(0)
        s.paragraph_format.space_before = Pt(10 if name != "Heading 1" else 14)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.keep_with_next = True


def add_body_paragraph(doc, text, compact=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Cm(0 if compact else 0.74)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    add_text_with_superscript_citations(p, text, size=9.5 if compact else 10.5)
    return p


def build():
    if not SOURCE.exists() or not REFS.exists():
        raise FileNotFoundError("Missing manuscript source or verified references")
    doc = Document()
    configure_styles(doc)
    doc.core_properties.title = "数据驱动揭示防腐底漆与聚氨酯/氟树脂面漆的构效关系及高性能配方区域"
    doc.core_properties.subject = "防腐涂层构效关系、机器学习与高性能配方区域"
    doc.core_properties.keywords = "防腐涂层; 环氧底漆; 聚氨酯/氟树脂面漆; 机器学习; 构效关系; 配方区域; 耐磨损"

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    first_title = True
    in_front_abstract = False
    abstract_label = None
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line in FIGURES:
            add_figure(doc, *FIGURES[line])
            continue
        if line == "[[PAGEBREAK]]":
            doc.add_page_break()
            continue
        if line in TABLES:
            TABLES[line](doc)
            continue
        if line == "[[REFERENCES]]":
            add_references(doc)
            continue
        if line.startswith("# "):
            text = normalize(line[2:])
            if first_title:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.keep_with_next = True
                set_run_font(p.add_run(text), east="黑体", size=14.5, bold=True)
                first_title = False
            else:
                doc.add_paragraph(text, style="Heading 1")
            in_front_abstract = False
            abstract_label = None
            continue
        if line.startswith("## "):
            text = normalize(line[3:])
            if text.startswith("Data-Driven Formulation-Property Relationships"):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(5)
                p.paragraph_format.keep_with_next = True
                set_run_font(p.add_run(text), size=11.0, bold=True)
            elif text in {"摘要", "Abstract"}:
                in_front_abstract = True
                abstract_label = text
            else:
                doc.add_paragraph(text, style="Heading 2")
            continue
        if re.match(r"^\d+\.\s", line) and line[0].isdigit():
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Cm(0.74)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            p.paragraph_format.space_after = Pt(3)
            add_text_with_superscript_citations(p, re.sub(r"^\d+\.\s*", "", line), size=10.5)
            continue
        if line.startswith("关键词：") or line.startswith("Key words:"):
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.space_after = Pt(5)
            label, value = line.split("：", 1) if "：" in line else line.split(":", 1)
            set_run_font(p.add_run(label + ("：" if "：" in line else ": ")), east="黑体" if "关键词" in label else "Times New Roman", size=9.5, bold=True)
            set_run_font(p.add_run(normalize(value.strip())), size=9.5)
            in_front_abstract = False
            abstract_label = None
            continue
        if line.startswith("Rui Zhou"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            authors = [("Rui Zhou", "a"), ("Luyao Bao", "a*"), ("Meirong Cai", "a*")]
            for index, (name, marker) in enumerate(authors):
                set_run_font(p.add_run(name), size=9.5)
                marker_run = p.add_run(marker)
                set_run_font(marker_run, size=9.5)
                marker_run.font.superscript = True
                if index < len(authors) - 1:
                    set_run_font(p.add_run(", "), size=9.5)
            continue
        if line.startswith(("a. State Key Laboratory", "*Corresponding author")):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            set_run_font(p.add_run(normalize(line)), size=9.5, italic=line.startswith("*Corresponding"))
            continue
        if in_front_abstract:
            if abstract_label:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.0
                label = abstract_label + ("：" if abstract_label == "摘要" else ": ")
                set_run_font(p.add_run(label), east="黑体" if abstract_label == "摘要" else "Times New Roman", size=9.5, bold=True)
                add_scientific_text(p, line, size=9.5)
                abstract_label = None
            else:
                add_body_paragraph(doc, line, compact=True)
        else:
            add_body_paragraph(doc, line)

    doc.save(OUTPUT)
    print(f"MANUSCRIPT_BUILT path={OUTPUT} bytes={OUTPUT.stat().st_size} paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")


if __name__ == "__main__":
    build()
